from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app import models
from app.user_access_report import user_access_workbook


def test_user_access_report_separates_effective_access_from_assigned_grants():
    engine = create_engine('sqlite:///:memory:')
    models.Base.metadata.create_all(engine)
    db = Session(engine)
    parent = models.QAWorkspace(workspace_key='PARENT', name='Parent QA', is_active=True)
    child = models.QAWorkspace(workspace_key='CHILD', name='Child QA', is_active=True,
                               parent_workspace=parent)
    admin = models.User(username='admin', full_name='Administrator', hashed_password='x',
                        is_active=True, role_assignments=[models.UserRole(role='ADMIN')])
    viewer = models.User(username='viewer', full_name='Parent Viewer', hashed_password='x',
                         is_active=True, department_assignments=[models.UserDepartment(department='QA')],
                         role_assignments=[models.UserRole(role='VIEW_ONLY')])
    disabled = models.User(username='disabled', full_name='Disabled User', hashed_password='x',
                           is_active=False, role_assignments=[models.UserRole(role='QA_ENGINEER')])
    db.add_all([parent, child, admin, viewer, disabled]); db.flush()
    db.add_all([
        models.QAWorkspaceMember(workspace_id=parent.id, user_id=viewer.id,
                                 role='PARENT_WORKSPACE_VIEWER', is_active=True),
        models.QAWorkspaceMember(workspace_id=child.id, user_id=disabled.id,
                                 role='WORKSPACE_MEMBER', is_active=True),
    ])
    db.commit()

    workbook = load_workbook(user_access_workbook(db, admin))
    assert workbook.sheetnames == [
        'Accounts and Roles', 'Effective Workspace Access',
        'Assigned Workspace Grants', 'Coordinator Scopes',
    ]
    accounts = list(workbook['Accounts and Roles'].values)[4:]
    viewer_account = next(row for row in accounts if row[1] == 'viewer')
    disabled_account = next(row for row in accounts if row[1] == 'disabled')
    assert viewer_account[8:10] == ('QA', 'VIEW_ONLY')
    assert viewer_account[13] == 2
    assert disabled_account[13] == 0

    access = list(workbook['Effective Workspace Access'].values)[4:]
    inherited = next(row for row in access if row[1] == 'viewer' and row[4] == 'Child QA')
    assert inherited[6:10] == ('Inherited parent grant', 'Read only', None, 'PARENT_WORKSPACE_VIEWER')
    assert not any(row[1] == 'disabled' for row in access)
    grants = list(workbook['Assigned Workspace Grants'].values)[4:]
    disabled_grant = next(row for row in grants if row[1] == 'disabled')
    assert disabled_grant[5:9] == ('WORKSPACE_MEMBER', 'Active', 'Active', 'No')
    db.close()
    engine.dispose()
