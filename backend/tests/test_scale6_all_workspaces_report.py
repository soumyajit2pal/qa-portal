import datetime
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.requests import Request
from starlette.responses import Response
from app import models
from app.constants import Role
from app.deps import _enforce_view_only_request, require_roles
from app.routers import auth, reports
from app.routers.export import _rows_to_xlsx, _rows_to_csv, _rows_to_pdf
from app.workspace_service import active_workspace_scope_ids, inherited_workspace_access_mode, selectable_workspace_ids, workspace_context


def request(method, path):
    return Request({'type': 'http', 'method': method, 'path': path, 'raw_path': path.encode(),
                    'query_string': b'', 'headers': [], 'scheme': 'http',
                    'server': ('test', 80), 'client': ('test', 1)})


@pytest.fixture
def setup():
    engine = create_engine('sqlite:///:memory:')
    models.Base.metadata.create_all(engine)
    with Session(engine) as db:
        parent = models.QAWorkspace(workspace_key='PARENT', name='Parent', is_active=True)
        child = models.QAWorkspace(workspace_key='CHILD', name='Child', is_active=True, parent_workspace=parent)
        other = models.QAWorkspace(workspace_key='OTHER', name='Other', is_active=True)
        inactive = models.QAWorkspace(workspace_key='INACTIVE', name='Inactive', is_active=False)
        user = models.User(username='scale.viewer', full_name='Scale Viewer', hashed_password='x',
                           is_active=True, role_assignments=[models.UserRole(role=Role.SCALE_6_PLUS)])
        db.add_all([parent, child, other, inactive, user]); db.commit()
        user.active_qa_workspace_id = parent.id
        user.active_workspace_scope_ids = tuple(active_workspace_scope_ids(db, user, parent.id))
        yield db, user, parent, child, other, inactive
    engine.dispose()


def test_scale6_selects_active_workspaces_without_granting_memberships(setup):
    db, user, parent, child, other, inactive = setup
    assert selectable_workspace_ids(db, user) == {parent.id, child.id, other.id}
    assert active_workspace_scope_ids(db, user, parent.id) == {parent.id, child.id}
    assert inherited_workspace_access_mode(db, user, other.id) == 'PARENT_VIEWER'
    out = auth.me(response=Response(), db=db, current_user=user)
    assert {item.workspace_id for item in out.workspace_access} == {parent.id, child.id, other.id}
    assert all(item.role == 'WORKSPACE_VIEWER' for item in out.workspace_access)
    assert not user.qa_workspace_access


def test_scale6_is_read_only_even_when_combined_with_workflow_or_admin_roles(setup):
    db, user, *_ = setup
    user.role_assignments.append(models.UserRole(role=Role.ADMIN))
    user.role_assignments.append(models.UserRole(role=Role.QA_ENGINEER)); db.commit()
    assert require_roles(Role.QA_ENGINEER)(request('GET', '/api/test-execution'), user) is user
    for path in ('/api/test-execution', '/api/qa-requests', '/api/workspaces', '/api/auth/users'):
        with pytest.raises(HTTPException) as error:
            _enforce_view_only_request(request('POST', path), user, db)
        assert error.value.status_code == 403
    _enforce_view_only_request(request('PATCH', '/api/workspaces/preference/current'), user, db)


def test_all_data_report_is_workspace_wise_and_counts_parent_once(setup):
    db, user, parent, child, other, inactive = setup
    created = datetime.datetime(2026, 9, 15, 10, 0)
    qa_request = models.QARequest(
        request_id='TQA-REQ-REPORT', application_name='CBS', department='IT - Software',
        qa_workspace_id=other.id, requester_id=user.id, status='RAISED',
        change_description='New payment flow', created_at=created,
    )
    project = models.TestProject(
        project_key='TQA-PROJ-REPORT', name='Payment', qa_workspace_id=other.id,
        department='IT - Software', created_by_id=user.id, owner_id=user.id, created_at=created,
    )
    db.add_all([qa_request, project]); db.flush()
    other_user = models.User(username='draft.owner', full_name='Draft Owner',
                             hashed_password='x', is_active=True,
                             role_assignments=[models.UserRole(role=Role.REQUESTER)])
    db.add(other_user); db.flush()
    private_draft = models.QARequest(
        application_name='Private Draft', department='IT - Software',
        qa_workspace_id=other.id, requester_id=other_user.id,
        status='DRAFT', created_at=created,
    )
    db.add(private_draft)
    db.add(models.FunctionalRequest(
        request_id='TQA-FUNC-REPORT', qa_request_id=qa_request.id,
        status='EXECUTION_IN_PROGRESS', created_at=created,
    ))
    cases = [models.TestCase(
        test_case_key=f'TQA-TC-REPORT-{number}', project_id=project.id,
        origin_workspace_id=other.id, test_scenario=f'Scenario {number}',
        status='Approved', is_deleted=False, created_at=created,
    ) for number in (1, 2)]
    cycle = models.TestCycle(
        cycle_key='TQA-CYCLE-REPORT', project_id=project.id, name='Regression',
        origin_workspace_id=other.id, status='In Progress', created_at=created,
    )
    db.add_all([*cases, cycle]); db.flush()
    executions = [models.TestExecution(
        cycle_id=cycle.id, test_case_id=case.id, status=status, created_at=created,
    ) for case, status in zip(cases, ('Pass', 'Fail'))]
    db.add_all(executions); db.flush()
    db.add_all([models.TestExecutionRun(
        execution_id=execution.id, attempt_no=1, status=execution.status,
        executed_at=created,
    ) for execution in executions])
    db.add(models.Defect(
        defect_key='TQA-DEF-REPORT', title='Failed payment', description='Failure',
        status='New', qa_workspace_id=other.id, cycle_id=cycle.id,
        application_name='CBS', module_feature='Payment', environment='UAT',
        severity='High', priority='P2', steps_to_reproduce='Run',
        expected_result='Pass', actual_result='Fail', reporter_id=user.id,
        reported_at=created, created_at=created,
    ))
    db.add(models.SuppressionRequest(
        suppression_id='TQA-SUP-REPORT', application_name='CBS',
        qa_workspace_id=other.id, department='IT - Software',
        scan_type='SAST', status='Security Approval Pending', created_at=created,
    ))
    db.commit()
    rows = reports.all_data_report(date_from='2026-09-01', db=db, current_user=user)
    def summary(module, status):
        return next(row for row in rows if row['Workspace'] == other.name
                    and row['Row Type'] == 'Summary' and row['Module'] == module
                    and row['Status'] == status)
    assert summary('QA Request', 'Total')['Count'] == 1
    assert summary('Functional Request', 'Total')['Count'] == 1
    assert summary('Testcase', 'Total')['Count'] == 2
    assert summary('Test Execution', 'Pass')['Count'] == 1
    assert summary('Test Execution', 'Fail')['Count'] == 1
    assert summary('Execution Attempt', 'Pass')['Count'] == 1
    assert summary('Execution Attempt', 'Fail')['Count'] == 1
    assert summary('Defect', 'New')['Count'] == 1
    assert summary('Defect Severity', 'High')['Count'] == 1
    assert summary('Suppression Request', 'Security Approval Pending')['Count'] == 1
    cycle_row = next(row for row in rows if row['Row Type'] == 'Record'
                     and row['Record ID'] == 'TQA-CYCLE-REPORT')
    assert cycle_row['Testcases'] == 2
    assert 'Pass 1' in cycle_row['Execution Results']
    assert 'Fail 1' in cycle_row['Execution Results']
    assert 'Attempts 2' in cycle_row['Execution Results']
    assert cycle_row['Defects'] == 1
    assert any(row['Record ID'] == 'TQA-REQ-REPORT' and row['Description'] == 'New payment flow'
               for row in rows)
    assert not any(row['Record ID'] == f'Draft #{private_draft.id}' for row in rows)
    assert not any(row['Workspace'] == inactive.name for row in rows)

    filtered_rows = reports.all_data_report(date_from='2026-09-16', db=db, current_user=user)
    assert all(row['Row Type'] == 'Summary' and row['Count'] == 0 for row in filtered_rows)

    user.role_assignments[:] = [row for row in user.role_assignments if row.role != Role.SCALE_6_PLUS]
    user.qa_workspace_memberships.append(models.QAWorkspaceMember(
        workspace_id=parent.id, role='PARENT_WORKSPACE_VIEWER', is_active=True,
    ))
    db.flush()
    with workspace_context(parent.id, (parent.id, child.id)):
        limited_rows = reports.all_data_report(db=db, current_user=user)
    assert not any(row['Record ID'] == 'TQA-REQ-REPORT' for row in limited_rows)
    assert not any(row['Record ID'] == 'TQA-CYCLE-REPORT' for row in limited_rows)

    user.qa_workspace_memberships.append(models.QAWorkspaceMember(
        workspace_id=other.id, role='WORKSPACE_MEMBER', is_active=True,
    ))
    db.flush()
    with workspace_context(parent.id, (parent.id, child.id)):
        multi_workspace_rows = reports.all_data_report(db=db, current_user=user)
    assert any(row['Record ID'] == 'TQA-REQ-REPORT' for row in multi_workspace_rows)
    assert any(row['Record ID'] == 'TQA-CYCLE-REPORT' for row in multi_workspace_rows)



def test_all_data_report_has_zero_totals_for_each_visible_workspace_on_empty_database(setup):
    db, user, *_ = setup
    rows = reports.all_data_report(db=db, current_user=user)
    assert len(rows) == 3 * len(reports._ALL_DATA_MODULES)
    assert all(row['Row Type'] == 'Summary' and row['Status'] == 'Total'
               and row['Count'] == 0 for row in rows)


def test_workspace_report_exports_organized_sections(setup):
    from openpyxl import load_workbook
    db, user, *_ = setup
    rows = reports.all_data_report(db=db, current_user=user)
    meta = {'report_name': 'All Data Report', 'module': 'all-data-report',
            'generated_at': '2026-09-15 IST', 'generated_by': user.full_name,
            'filters': '', 'total_records': len(rows)}
    wb = load_workbook(_rows_to_xlsx(rows, meta))
    assert wb.sheetnames == [
        'Guide', 'Workspace Overview', 'Status Breakdown', 'Requests', 'Projects',
        'Testcases', 'Cycles', 'Executions', 'Attempts', 'Defects', 'Governance',
    ]
    assert wb['Workspace Overview']['A8'].value == 'Workspace'
    assert wb['Workspace Overview']['B8'].value == 'Main QA Requests'
    assert wb['Workspace Overview']['B9'].value == 0
    assert wb['Workspace Overview'].freeze_panes == 'B9'
    assert wb['Requests']['A9'].value == 'No records found'
    csv_output = _rows_to_csv(rows, meta).getvalue()
    assert 'Workspace Overview,3 rows' in csv_output
    assert 'Main QA Requests' in csv_output
    assert 'Requests,0 rows' in csv_output
    assert _rows_to_pdf(rows, meta).getvalue().startswith(b'%PDF')


def test_workspace_report_overview_and_details_use_existing_counts():
    from openpyxl import load_workbook
    from app.workspace_report_layout import all_data_sections
    schema = {
        'Workspace': 'Quality', 'Record ID': '', 'Parent / Linked ID': '',
        'Description': '', 'Application / Project': '', 'Department': '',
        'Testcases': '', 'Execution Results': '', 'Defects': '', 'Created At': '',
    }
    def summary(module, status, count):
        return {**schema, 'Row Type': 'Summary', 'Module': module, 'Status': status, 'Count': count}
    def record(module, key, status, **fields):
        return {**schema, 'Row Type': 'Record', 'Module': module,
                'Record ID': key, 'Status': status, 'Count': 1, **fields}
    rows = [
        summary('QA Request', 'Total', 1), summary('QA Request', 'RAISED', 1),
        summary('Functional Request', 'Total', 1), summary('Functional Request', 'EXECUTION_IN_PROGRESS', 1),
        summary('Test Execution', 'Total', 2), summary('Test Execution', 'Pass', 1),
        summary('Test Execution', 'Fail', 1), summary('Execution Attempt', 'Total', 3),
        summary('Defect', 'Total', 1), summary('Defect', 'New', 1),
        summary('Defect Severity', 'Total', 1), summary('Defect Severity', 'High', 1),
        record('QA Request', 'TQA-REQ-1', 'RAISED', Description='Payment flow'),
        record('Functional Request', 'TQA-FUNC-1', 'EXECUTION_IN_PROGRESS',
               **{'Parent / Linked ID': 'TQA-REQ-1'}),
        record('Defect', 'TQA-DEF-1', 'New', **{'Execution Results': 'Severity High'}),
    ]
    sections = {title: values for title, values, _ in all_data_sections(rows)}
    overview = sections['Workspace Overview'][0]
    assert (overview['Main QA Requests'], overview['Functional'], overview['Execution Slots'],
            overview['Passed'], overview['Failed'], overview['Attempts'], overview['Defects'],
            overview['High Defects']) == (1, 1, 2, 1, 1, 3, 1, 1)
    assert sections['Requests'][1]['Parent QA Request'] == 'TQA-REQ-1'
    assert sections['Defects'][0]['Severity'] == 'High'
    assert not any(item['Status'] == 'Total' for item in sections['Status Breakdown'])
    meta = {'report_name': 'All Data Report', 'module': 'all-data-report',
            'generated_at': '2026-09-15 IST', 'generated_by': 'Tester',
            'filters': '', 'total_records': 3}
    wb = load_workbook(_rows_to_xlsx(rows, meta))
    assert wb['Requests']['C9'].value == 'TQA-REQ-1'
    assert wb['Defects']['G9'].value == 'High'
